import asyncio
import httpx
import json as Json
import os
import base64
import argparse
import io

from docx import Document
from importlib.resources import read_text
from datetime import datetime
from sanic import Sanic
from sanic.response import json
from prometheus_client import generate_latest, CONTENT_TYPE_LATEST
from app.config import Config
from app.logger import setup_logger
from app.telemetry import REQUEST_COUNT, REQUEST_LATENCY, record_request

#Configs
OLLAMA_DOCKER= "http://ollama:11434"
OLLAMA_GENERAL= "http://localhost:11434"
MODEL = "phi4:latest"  # Consistent model name
TIMEOUT=1440

# Parse command-line arguments
parser = argparse.ArgumentParser(description="Run with Docker or Local Ollama API.")
parser.add_argument(
    "--docker",
    action="store_true",
    help="Use the Docker URL (http://ollama:11434) instead of localhost.",
)
args = parser.parse_args()

# Determine OLLAMA_URL based on the argument
OLLAMA_URL = f"{OLLAMA_DOCKER if args.docker else OLLAMA_GENERAL}/api/generate"

# Initialize Sanic app
app = Sanic("SanicServer", strict_slashes=True)
config = Config()
logger = setup_logger(config.LOG_FILE)

#Load all the prompts 
with open(os.path.join(os.getcwd(), "promptrepo", "prompt.json"), "r") as file:
    data = Json.load(file)

templatefillprompt= data.get("templatefill", "")
comparedocsprompt= data.get("comparedocs", "")
docretrivalprompt= data.get("docretrival", "")
summarizedocsprompt= data.get("summarizedocs", "")
trustreconprompt= data.get("trustrecon", "")

@app.middleware("request")
async def before_request(request):
    record_request(request.method, request.path)

@app.middleware("response")
async def after_request(request, response):
    REQUEST_LATENCY.set(response.elapsed.total_seconds() if hasattr(response, "elapsed") else 0)

@app.route("/")
async def home(request):
    logger.info("Home endpoint accessed")
    return json({"message": "Welcome to Sanic!"})

@app.get("/metrics")
async def metrics(request):
    return app.response.raw(
        generate_latest(),
        headers={"Content-Type": CONTENT_TYPE_LATEST},
    )

@app.route("/test", methods=["POST", "PUT"])
async def handler(request):
    return json({"message": "Test endpoint works!"})

async def check_ollama_server():
    """Check if Ollama server is available"""
    try:
        async with httpx.AsyncClient() as client:
            response = await client.get(OLLAMA_URL.split('/api')[0])
            return response.status_code == 200
    except Exception:
        return False

async def handle_ollama_request(payload):
    """Handle Ollama API request with proper error handling"""
    if not await check_ollama_server():
        raise ConnectionError(f"Ollama server is not available. Please ensure it's running on {OLLAMA_URL}")
        
    async with httpx.AsyncClient(timeout=TIMEOUT) as client:
        try:
            response = await client.post(OLLAMA_URL, json=payload)
            response.raise_for_status()
            return response.json()
        except httpx.HTTPStatusError as e:
            error_message = f"HTTP Error {e.response.status_code}"
            try:
                error_detail = e.response.json()
                error_message += f": {error_detail.get('error', '')}"
            except Exception:
                error_message += f": {e.response.read().decode('utf-8', 'ignore')}"
            raise httpx.HTTPStatusError(error_message, request=e.request, response=e.response)

async def docstringtotext(inputdocstring):
        base64_data=inputdocstring.split(",")[1]

        # Step 2: Decode the Base64 data
        docx_bytes = base64.b64decode(base64_data)

        # Step 3: Load the document from the decoded bytes
        docx_file = io.BytesIO(docx_bytes)
        document = Document(docx_file)

        # Step 4: Extract text from the document
        full_text = []
        for para in document.paragraphs:
            full_text.append(para.text)

        # Combine all paragraphs into a single string
        realtext = '\n'.join(full_text)
        return realtext

async def generatedocstring(inputresponsestring):
        doc = Document()
        for line in inputresponsestring.split('\n'):
            doc.add_paragraph(line.strip())
        #response_data
        docx_buffer = io.BytesIO()
        doc.save(docx_buffer)
        docx_buffer.seek(0)

        encoded_base64 = base64.b64encode(docx_buffer.read()).decode('utf-8')
        # Create the base64 data URI
        returndata_uri = f"data:@file/vnd.openxmlformats-officedocument.wordprocessingml.document;base64,{encoded_base64}"
        return returndata_uri

@app.post("/templatefill")
async def templatefill(request):
    """Route to proxy requests to Ollama for template fill behaviour"""
    try:
        logger.info("Ollama endpoint accessed.")
        
        # Extract data from the incoming request correctly
        data = request.json
        info = data.get("InputInfo", "")
        template = data.get("InputTemplate", "")
        instruction = data.get("InputInstruction", "")
        input_template_text =  await docstringtotext(template)
        
        # Ensure templatefill is not None
        if not templatefillprompt:
            raise ValueError("Templatefill is not loaded correctly")

        # Format the template with the request data
        prompt = templatefillprompt.format(input_info=info, input_template=input_template_text, input_instruction=instruction)

        logger.debug(f"Formatted prompt: {prompt}")  # Log the formatted prompt

        payload = {
            "model": MODEL,  
            "prompt": prompt,  # Use the formatted prompt here
            "stream": False
        }
        
        response_data = await handle_ollama_request(payload)
        returndocx = await generatedocstring(response_data.get("response", "").strip())

        return json({
            "response": returndocx,
            "tokens_used": response_data.get("total_tokens", 0)
        })

    except ValueError as e:
        logger.error(f"Invalid Request: {str(e)}")
        return json({"error": "Invalid JSON request"}, status=400)
        
    except ConnectionError as e:
        logger.error(f"Connection Error: {str(e)}")
        return json({"error": str(e)}, status=503)

    except httpx.HTTPStatusError as e:
        logger.error(f"HTTP Error: {e.response.status_code} - {e.response.text}")
        return json({"error": f"HTTP Error: {e.response.status_code}"}, status=e.response.status_code)

    except httpx.RequestError as e:
        logger.error(f"Request Error: {str(e)}")
        return json({"error": "Failed to communicate with Ollama"}, status=500)

    except Exception as e:
        logger.error(f"Unexpected Error: {str(e)}")
        return json({"error": "An unexpected error occurred"}, status=500)

@app.post("/summarizedocs")
async def summarizedocs(request):
    """Route to proxy requests to Ollama for summarize docs behaviour"""
    try:
        logger.info("Ollama endpoint accessed.")
        
        # Extract data from the incoming request correctly
        data = request.json
        text = data.get("Input", "")

        input_doc_text =  await docstringtotext(text)

        # Ensure summarizedocs prompt is not None
        if not summarizedocsprompt:
            raise ValueError("Summarizedocs prompt is not loaded correctly")

        # Format the template with the request data
        prompt = summarizedocsprompt.format(input_text=input_doc_text)

        logger.debug(f"Formatted prompt: {prompt}")  # Log the formatted prompt

        payload = {
            "model": MODEL,  
            "prompt": prompt,  # Use the formatted prompt here
            "stream": False  
        }
        
        response_data = await handle_ollama_request(payload)

        returndocx = await generatedocstring(response_data.get("response", "").strip())

        return json({
            "response": returndocx,
            "tokens_used": response_data.get("total_tokens", 0)
        })

    except ValueError as e:
        logger.error(f"Invalid Request: {str(e)}")
        return json({"error": "Invalid JSON request"}, status=400)
        
    except ConnectionError as e:
        logger.error(f"Connection Error: {str(e)}")
        return json({"error": str(e)}, status=503)

    except httpx.HTTPStatusError as e:
        logger.error(f"HTTP Error: {str(e)}")
        return json({"error": str(e)}, status=e.response.status_code)

    except httpx.RequestError as e:
        logger.error(f"Request Error: {str(e)}")
        return json({"error": "Failed to communicate with Ollama"}, status=500)

    except Exception as e:
        logger.error(f"Unexpected Error: {str(e)}")
        return json({"error": "An unexpected error occurred"}, status=500)

@app.post("/comparedocs")
async def comparedocs(request):
    """Route to proxy requests to Ollama for compare docs behaviour"""
    try:
        logger.info("Ollama endpoint accessed.")
        
        # Extract data from the incoming request correctly
        data = request.json
        docuemntA = data.get("Doc_A", "")
        docuemntB = data.get("Doc_B", "")

        input_doc_A =  await docstringtotext(docuemntA)
        input_doc_B =  await docstringtotext(docuemntB)

        # Ensure comparedocs prompt is not None
        if not comparedocsprompt:
            raise ValueError("Comparedocs prompt is not loaded correctly")

        # Format the template with the request data
        prompt = comparedocsprompt.format(doc_a=input_doc_A, doc_b=input_doc_B)
        logger.debug(f"Formatted prompt: {prompt}")  # Log the formatted prompt

        payload = {
            "model": MODEL,  
            "prompt": prompt,  # Use the formatted prompt here
            "stream": False  
        }
        
        response_data = await handle_ollama_request(payload)

        returndocx = await generatedocstring(response_data.get("response", "").strip())

        return json({
            "response": returndocx,
            "tokens_used": response_data.get("total_tokens", 0)
        })

    except ValueError as e:
        logger.error(f"Invalid Request: {str(e)}")
        return json({"error": "Invalid JSON request"}, status=400)
        
    except ConnectionError as e:
        logger.error(f"Connection Error: {str(e)}")
        return json({"error": str(e)}, status=503)

    except httpx.HTTPStatusError as e:
        logger.error(f"HTTP Error: {e.response.status_code} - {e.response.text}")
        return json({"error": f"HTTP Error: {e.response.status_code}"}, status=e.response.status_code)

    except httpx.RequestError as e:
        logger.error(f"Request Error: {str(e)}")
        return json({"error": "Failed to communicate with Ollama"}, status=500)

    except Exception as e:
        logger.error(f"Unexpected Error: {str(e)}")
        return json({"error": "An unexpected error occurred"}, status=500)

@app.post("/docretrival")
async def docretrival(request):
    """Route to proxy requests to Ollama for doc retrival behaviour"""
    try:
        logger.info("Ollama endpoint accessed.")
        
        # Extract data from the incoming request correctly
        data = request.json
        text = data.get("Input", "")

        input_doc_text =  await docstringtotext(text)

        # Ensure templatefill is not None
        if not templatefillprompt:
            raise ValueError("Templatefill is not loaded correctly")


        # Ensure docretrival prompt is not None
        if not docretrivalprompt:
            raise ValueError("Docretrival prompt is not loaded correctly")

        # Format the template with the request data
        prompt = docretrivalprompt.format(input_text=input_doc_text)
        logger.debug(f"Formatted prompt: {prompt}")  # Log the formatted prompt

        payload = {
            "model": MODEL,  
            "prompt": prompt,  # Use the formatted prompt here
            "stream": False  
        }
        
        response_data = await handle_ollama_request(payload)

        returndocx = await generatedocstring(response_data.get("response", "").strip())


        return json({
            "response": returndocx ,
            "tokens_used": response_data.get("total_tokens", 0)
        })

    except ValueError as e:
        logger.error(f"Invalid Request: {str(e)}")
        return json({"error": "Invalid JSON request"}, status=400)
        
    except ConnectionError as e:
        logger.error(f"Connection Error: {str(e)}")
        return json({"error": str(e)}, status=503)

    except httpx.HTTPStatusError as e:
        logger.error(f"HTTP Error: {e.response.status_code} - {e.response.text}")
        return json({"error": f"HTTP Error: {e.response.status_code}"}, status=e.response.status_code)

    except httpx.RequestError as e:
        logger.error(f"Request Error: {str(e)}")
        return json({"error": "Failed to communicate with Ollama"}, status=500)

    except Exception as e:
        logger.error(f"Unexpected Error: {str(e)}")
        return json({"error": "An unexpected error occurred"}, status=500)


@app.post("/trustrecon")
async def trustrecon(request):
  return json({"message": "trustrecon endpoint works!"})



if __name__ == "__main__":
    print("Registered Routes:", list(app.router.routes_all.keys()))
    ollama_status = asyncio.run(check_ollama_server())
    print(f"Ollama Connection :-  {ollama_status}")
    print(config.LOG_FILE)
    app.run(host=config.HOST, port=config.PORT, debug=config.DEBUG)
